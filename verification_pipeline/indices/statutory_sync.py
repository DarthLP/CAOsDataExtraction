"""
STATUTORY SYNC — statutory_timeline.xlsx is the GROUND TRUTH (Hanna edits it in Excel);
this script derives the machine files the indices read:

    statutory_timeline.xlsx  --sync-->  statutory_all.csv + wml_timeline.csv

History: build_statutory.py used to GENERATE the workbook from its DATA list, which
silently overwrote Hanna's manual research (recovered 2026-07-06 from
_recovered/statutory_timeline.BACKUP8.xlsx). Direction is now reversed:
NEVER regenerate the workbook; edit the workbook, then run this sync.

Parsing per topic tab (all tabs except Overview): the row containing 'Variable' is the
header; columns = Variable, From, To, Value, Unit, Direction(role), Statutory law,
Source 1, Source 2, Notes, field (machine). Rows without a machine field are still
exported (documentation) but only rows WITH a field and a numeric value can impute.
Wage tab additionally holds the WML series block (sub-header 'Valid from') ->
wml_timeline.csv.

Run: python3.13 indices/statutory_sync.py   (then rerun the index pipeline)
"""
import os, csv, sys
from datetime import datetime, date
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import index_lib as il

XLSX = il.locate("statutory_timeline.xlsx")   # Hanna's hand-maintained ground truth -> review/


def iso(v):
    if v is None: return ""
    if isinstance(v, (datetime, date)): return v.strftime("%Y-%m-%d")
    return str(v).strip()


def cell_link(cell):
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    return str(cell.value).strip() if cell.value is not None else ""


def main():
    wb = load_workbook(XLSX)
    rows, wml = [], []
    for tab in wb.sheetnames:
        if tab == "Overview": continue
        ws = wb[tab]
        hdr = None
        for r in range(1, 9):
            if any(ws.cell(row=r, column=c).value == "Variable" for c in range(1, 4)):
                hdr = r; break
        if hdr is None: continue
        in_series = False
        for r in range(hdr + 1, ws.max_row + 1):
            c1 = iso(ws.cell(row=r, column=1).value)
            if not c1: continue
            if tab == "Wage" and c1 == "Valid from":
                in_series = True; continue
            if tab == "Wage" and in_series:
                # WML series row: Valid from | Valid to | EUR/month | EUR/hour | basis | source
                month = ws.cell(row=r, column=3).value
                hour = ws.cell(row=r, column=4).value
                basis = iso(ws.cell(row=r, column=5).value)
                src = cell_link(ws.cell(row=r, column=6))
                if month is None and hour is None: continue
                ver = "YES" if "VERIFY" in (basis + src).upper() else ""
                wml.append([c1, month if month is not None else "",
                            hour if hour is not None else "", basis, ver, src])
                continue
            var = c1
            if var.startswith("Note") or var.startswith("Statutory minimum wage — full"): continue
            role = iso(ws.cell(row=r, column=6).value)
            rows.append([
                tab.lower(),                                   # topic
                iso(ws.cell(row=r, column=11).value),          # field (machine)
                role,
                iso(ws.cell(row=r, column=2).value),           # effective_from
                iso(ws.cell(row=r, column=3).value),           # effective_to
                iso(ws.cell(row=r, column=4).value),           # value
                iso(ws.cell(row=r, column=5).value),           # unit
                iso(ws.cell(row=r, column=7).value),           # legal_basis
                cell_link(ws.cell(row=r, column=8)),           # source (Source 1)
                (iso(ws.cell(row=r, column=10).value) +
                 (f" [source2: {cell_link(ws.cell(row=r, column=9))}]"
                  if ws.cell(row=r, column=9).value else "")), # notes (+ Source 2)
            ])

    with open(os.path.join(il.OUT, "statutory_all.csv"), "w", newline="") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(["topic", "field", "statutory_role", "effective_from", "effective_to",
                    "value", "unit", "legal_basis", "source", "notes"])
        w.writerows(rows)
    with open(os.path.join(il.OUT, "wml_timeline.csv"), "w", newline="") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(["effective_from", "wml_month_eur", "wml_hour_eur", "basis", "please_verify", "source"])
        w.writerows(wml)
    n_field = sum(1 for r in rows if r[1])
    print(f"synced {len(rows)} statutory rows ({n_field} with machine field) + "
          f"{len(wml)} WML revisions from {os.path.basename(XLSX)}")


if __name__ == "__main__":
    main()
